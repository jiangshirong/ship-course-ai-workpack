"""Focused regression checks for portable helpers, not engineering certification."""
import copy,json,math,tempfile,unittest
from pathlib import Path
from course_formula import Workbook,FormulaError
from recalculate import load,evaluate,write_cached_xlsx
from check_geometry import check
from prepare_t_model import prepare
from select_profiles import combined_section,select
from load_cases import validate_case
from validate_xlsx import validate
from check_t_assignment import check as check_assignment
from verify_examples import verify
from check_delivery import check as check_delivery

ROOT=Path(__file__).resolve().parents[1]

class FormulaTests(unittest.TestCase):
    def setUp(self):self.w=Workbook({'Main':{'A1':-5,'A2':'header','B1':7},'Другой':{'A1':9},'Sheet2':{'A1':11}})
    def test_semantics(self):
        for f,expected in {'"foo"="bar"':False,'"Foo"="foo"':True,'IF(1,2,1/0)':2,'ROUND(2.5,0)':3,
            'ROUND(-2.5,0)':-3,'-2^2':4,'2^3^2':64,'Sheet2!A1':11,"'Другой'!A1":9,'Другой!A1':9,
            'SUM(A1:A2)':-5,'MAX(A1:A2)':-5,'COUNT(A1:A2)':1,'IF(FALSE,1/0,3)':3,
            'CEILING.MATH(-2.5,1)':-2,'CEILING.MATH(-2.5,1,1)':-3,'CEILING(2.1,1)':3,
            'IF(1,"Выполнено","Не выполнено")':'Выполнено','SUM(2,3^2)':11}.items():
            with self.subTest(formula=f):self.assertEqual(self.w.eval(f,'Main'),expected)
    def test_fail_closed(self):
        for f in ['A2+1','UNKNOWN(2)','1/0','SUM(Sheet2!A1:Main!B1)','"2"=2','ROUND(2)','[old.xlsx]Sheet1!A1']:
            with self.subTest(formula=f),self.assertRaises(Exception):self.w.eval(f,'Main')
        self.w.set('Main','B1','=B1+1')
        with self.assertRaises(FormulaError):self.w.value('Main','B1')
    def test_changed_input_invalidates_cache(self):
        self.w.set('Main','B2','=B1*2');self.assertEqual(self.w.value('Main','B2'),14)
        self.w.set('Main','B1',9);self.assertEqual(self.w.value('Main','B2'),18)
    def test_historical_formula_regression(self):
        for name in ['T/workbook.json','T/corrected_section_workbook.json','C_longitudinal/workbook.json']:
            path=ROOT/'templates'/name;spec=json.loads(path.read_text('utf-8'));book=Workbook.from_template(str(path));n=0
            for s in spec['sheets']:
                for c in s['cells']:
                    if not isinstance(c['value'],str)or not c['value'].startswith('='):continue
                    actual=book.value(s['name'],c['address']);expected=c.get('cached')
                    if expected is None:continue
                    with self.subTest(template=name,sheet=s['name'],cell=c['address']):
                        if isinstance(expected,(int,float))and not isinstance(expected,bool):self.assertTrue(math.isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-9),(actual,expected))
                        else:self.assertEqual(actual,expected)
                    n+=1
            self.assertGreater(n,100)
            print(f'Historical regression {name}: {n} cached formulas compared (not proof of engineering correctness)')
    def test_cache_roundtrip(self):
        from openpyxl import Workbook as X,load_workbook
        with tempfile.TemporaryDirectory() as d:
            source=Path(d)/'in.xlsx';out=Path(d)/'out.xlsx';x=X();s=x.active
            s['A1']=2;s['B1']='=A1+3';s['C1']='=IF(B1=5,"Готово","Ошибка")';s['D1']='=B1=5';x.save(source)
            original=source.read_bytes();write_cached_xlsx(source,out,evaluate(load(source)))
            self.assertEqual(source.read_bytes(),original)
            got=load_workbook(out,data_only=True).active
            self.assertEqual([got[a].value for a in ['B1','C1','D1']],[5,'Готово',True])
            self.assertEqual(load_workbook(out).active['B1'].value,'=A1+3')
            self.assertTrue(validate(out)['ok'])
            self.assertFalse(validate(source)['ok'])
            with self.assertRaises(FileExistsError):write_cached_xlsx(source,out,{})

class GeometryTests(unittest.TestCase):
    def setUp(self):self.cfg=json.loads((ROOT/'examples/t_model.example.json').read_text('utf-8'))
    def test_closure_and_input_failure(self):
        self.assertTrue(check(self.cfg['layout'])['ok'])
        for change in ['negative','arc','height','spacing','span']:
            g=copy.deepcopy(self.cfg['layout'])
            if change=='negative':g['flat_bands_m'][1]=-1
            elif change=='arc':g['arc_bands_m'][0]+=.1
            elif change=='height':g['workbook_inner_side_top_m']+=.38
            elif change=='spacing':g['longitudinals']['deck']['calculation_spacing_m']=.5
            else:g['longitudinals']['deck']['calculation_span_m']=3
            with self.subTest(change=change):self.assertFalse(check(g)['ok'])
    def test_prepared_model_links(self):
        spec=prepare(self.cfg);book=Workbook({s['name']:{c['address']:c['value']for c in s['cells']}for s in spec['sheets']})
        self.assertAlmostEqual(book.value('Балластные случаи','B9'),14.18)
        self.assertEqual(book.value('Балки','E22'),38);self.assertEqual(book.value('Балки','E23'),18)
        self.assertEqual(book.value('Эквивалентный брус','B96'),book.value('Сечения','E11'))
        book.set('Сечения','E11',99);self.assertEqual(book.value('Эквивалентный брус','B96'),99)
        for B,D,R in [(23.9,14.1,1.4),(27,15.5,1.55),(30,18,1.8)]:
            book.set('Исходные данные','B9',B);book.set('Исходные данные','B10',D);book.set('Исходные данные','B16',R)
            widths=[book.value('Эквивалентный брус',f'D{r}')for r in range(12,30)]
            self.assertGreater(min(widths),0)
            self.assertAlmostEqual(sum(widths),(B/2-R+math.pi*R/2+D-R)*1000)
            zi=book.value('Балластные случаи','B9');h=book.value('Исходные данные','B14')
            self.assertAlmostEqual(sum(book.value('Эквивалентный брус',f'D{r}')for r in range(33,40)),(zi-h)*1000)
        evaluate(book)
    def test_combined_section_against_two_rectangles(self):
        p={'h':100,'A':10,'y':5,'I':1*10**3/12}
        s=combined_section(p,100,10)
        # Exact rectangle integration about baseline, then shift to centroid.
        a=20;centroid=(10*.5+10*6)/a
        i0=10*1**3/3+(11**3-1**3)/3
        expected=(i0-a*centroid**2)/(11-centroid)
        self.assertAlmostEqual(s['W_cm3'],expected)

    def test_all_sixteen_groups_and_cl_binding(self):
        spec=prepare(self.cfg);book=Workbook({s['name']:{c['address']:c['value']for c in s['cells']}for s in spec['sheets']})
        catalog=json.loads((ROOT/'data/hp_catalog.json').read_text('utf-8'))['profiles']
        selection=select(book,catalog)
        self.assertEqual(len(selection),16);self.assertIn('CL-F4',selection)
        self.assertTrue(all(selection[f'CL-F{i}']['selected']for i in range(1,5)))
        book.set('Дополнительные связи','D37',1e9)
        result=check_assignment(book)
        self.assertTrue(any('CL-F1 W' in issue for issue in result['issues']))
        old=book.value('Дополнительные связи','F48')
        book.set('Параметры модели','G170',10000)
        self.assertGreater(book.value('Дополнительные связи','F48'),old)

    def test_pairing_and_explicit_assumptions(self):
        c=copy.deepcopy(self.cfg['deck_cases']['cargo'][0]);validate_case(c)
        for mode in ['draft','phase','envelope']:
            v=copy.deepcopy(c)
            if mode=='draft':v['outside_state']['draft_m']=5
            elif mode=='phase':v['outside_state']['phase_id']='other'
            else:v['inside_kind']='envelope'
            with self.subTest(mode=mode),self.assertRaises(ValueError):validate_case(v)
        c.update(combination_mode='no_counterpressure',outside_kpa=0,inside_kind='envelope');validate_case(c)
        cfg=copy.deepcopy(self.cfg);del cfg['design_inputs']['tank_length_m']
        with self.assertRaises(ValueError):prepare(cfg)
        cfg=copy.deepcopy(self.cfg);del cfg['profiles']['CL-F4']
        with self.assertRaises(ValueError):prepare(cfg)

    def test_known_arithmetic_answers(self):self.assertEqual(verify(),[])

    def test_delivery_format_contract(self):
        with tempfile.TemporaryDirectory() as d:
            folder=Path(d)/'Тест';folder.mkdir()
            (folder/'Тест_проект.dwg').write_bytes(b'dwg')
            (folder/'Тест_расчёты.xlsx').write_bytes(b'xlsx')
            self.assertTrue(check_delivery(folder,'Тест')['ok'])
            (folder/'Тест_проект.dxf').write_bytes(b'dxf')
            self.assertFalse(check_delivery(folder,'Тест')['ok'])

if __name__=='__main__':unittest.main(verbosity=2)
