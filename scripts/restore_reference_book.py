"""Restore an ANONYMIZED HISTORICAL reference, not a new student's assignment.
Does not evaluate formulas. Cached historical values are intentionally not copied.
"""
import argparse,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment
from openpyxl.workbook.properties import CalcProperties

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("spec",type=Path);p.add_argument("output",type=Path)
    args=p.parse_args()
    if args.output.suffix.lower()!=".xlsx":p.error("Output must be .xlsx")
    if args.output.exists():p.error("Output exists; choose a new path")
    spec=json.loads(args.spec.read_text("utf-8"))
    if spec.get("format")!="course-workbook-reference-v1":p.error("Unsupported reference format")
    w=Workbook();w.remove(w.active)
    w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
    for s in spec["sheets"]:
        ws=w.create_sheet(s["name"])
        for item in s["cells"]:
            c=ws[item["address"]];c.value=item["value"]
            c.number_format=item.get("number_format","General")
            c.font=Font(name="Arial",size=11)
            c.alignment=Alignment(vertical="top",wrap_text=True)
        for m in s["merged_ranges"]:ws.merge_cells(m)
        for key,width in s["column_widths"].items():
            if width:ws.column_dimensions[key].width=width
    w.properties.title="Historical reference — requires new inputs and recalculation"
    w.properties.creator="Course work pack"
    args.output.parent.mkdir(parents=True,exist_ok=True);w.save(args.output)
    print("Historical reference restored. Formulas NOT evaluated; this is NOT a completed assignment.")
if __name__=="__main__":main()
